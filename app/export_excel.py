import io

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models import InvoiceRecord


INVOICE_COLUMNS = [
    ("supplier_name", "Supplier"),
    ("invoice_number", "Invoice #"),
    ("invoice_date", "Date"),
    ("currency", "Currency"),
    ("subtotal", "Subtotal"),
    ("tax", "Tax"),
    ("total", "Total"),
]


def records_to_xlsx(
    records: list[InvoiceRecord],
    *,
    include_line_items: bool = True,
) -> bytes:
    """Build an .xlsx workbook with invoice rows and (optionally) line items."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"

    # Header
    for col_idx, (_, label) in enumerate(INVOICE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = Font(bold=True)

    # Rows
    for row_idx, rec in enumerate(records, start=2):
        ws.cell(row=row_idx, column=1, value=rec.supplier_name or "")
        ws.cell(row=row_idx, column=2, value=rec.invoice_number or "")
        ws.cell(row=row_idx, column=3, value=rec.invoice_date or "")
        ws.cell(row=row_idx, column=4, value=rec.currency or "")
        ws.cell(row=row_idx, column=5, value=rec.subtotal)
        ws.cell(row=row_idx, column=6, value=rec.tax)
        ws.cell(row=row_idx, column=7, value=rec.total)

    # Column widths
    for col_idx in range(1, len(INVOICE_COLUMNS) + 1):
        ws.column_dimensions[chr(64 + col_idx)].width = 18

    # Line items sheet
    if include_line_items and records:
        ws2 = wb.create_sheet("Line items")
        headers = ["Invoice #", "Description", "Qty", "Unit price", "Total"]
        for col_idx, label in enumerate(headers, start=1):
            cell = ws2.cell(row=1, column=col_idx, value=label)
            cell.font = Font(bold=True)

        row_idx = 2
        for rec in records:
            for li in rec.line_items:
                ws2.cell(row=row_idx, column=1, value=rec.invoice_number or "")
                ws2.cell(row=row_idx, column=2, value=li.description or "")
                ws2.cell(row=row_idx, column=3, value=li.quantity)
                ws2.cell(row=row_idx, column=4, value=li.unit_price)
                ws2.cell(row=row_idx, column=5, value=li.total)
                row_idx += 1

        for col_idx in range(1, 6):
            ws2.column_dimensions[chr(64 + col_idx)].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()